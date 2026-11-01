"""
qbo_ar_excel.py
===============
Converts a QBO AgedReceivableDetail JSON response into a formatted .xlsx file.

Layout
------
Rows 1-3  : Title block — company name, client name, as-of date
Row  4    : blank
Row  5    : Column header (bold, dark-navy fill, frozen below)
Per bucket: label row (bold, light-blue fill)
            data rows
            subtotal row (bold, medium-blue fill)
            blank separator
Last row  : Grand Total (bold, medium border-top)

Columns   : whatever QBO returns (typically Date, Type, Num, Due Date,
            Past Due, Amount, Open Balance).
Dates     : real Excel dates — MM/DD/YYYY format
Money     : real floats — #,##0.00 with thousands separator
File name : AR_Aging_{ClientName}_{YYYY-MM-DD}.xlsx
"""

import io
import logging
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ── Styles ────────────────────────────────────────────────────────────────────
_HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")   # dark navy
_BUCKET_FILL  = PatternFill("solid", fgColor="D6E4F0")   # light blue
_SUBTOT_FILL  = PatternFill("solid", fgColor="BDD7EE")   # medium blue
_GTOTAL_FILL  = PatternFill("solid", fgColor="9DC3E6")   # stronger blue

_HEADER_FONT  = Font(bold=True, color="FFFFFF", size=10)
_BUCKET_FONT  = Font(bold=True, size=10)
_BOLD_FONT    = Font(bold=True, size=10)
_NORMAL_FONT  = Font(size=10)
_TOP_BORDER   = Border(top=Side(style="medium"))

_DATE_FMT   = "MM/DD/YYYY"
_MONEY_FMT  = '#,##0.00'
_INT_FMT    = "#,##0"

# Map QBO ColType to how we format the Excel cell
_DATE_COLTYPES  = {"tx_date", "due_date"}
_MONEY_COLTYPES = {"subt_nat_amount", "rbal_nat_amount"}
_INT_COLTYPES   = {"aging"}   # "Past Due" (days)


# ── Public helpers ────────────────────────────────────────────────────────────

def has_data(report: dict) -> bool:
    """Return False when the report has no transaction rows (no open balance)."""
    if report.get("Header", {}).get("NoReportData"):
        return False
    for section in report.get("Rows", {}).get("Row", []):
        if section.get("Rows", {}).get("Row"):
            return True
    return False


def xlsx_filename(client_name: str) -> str:
    today = date.today().strftime("%Y-%m-%d")
    safe  = re.sub(r'[\\/:*?"<>|,]', "_", client_name).strip()
    return f"AR_Aging_{safe}_{today}.xlsx"


# ── Main builder ──────────────────────────────────────────────────────────────

def build_ar_aging_xlsx(report: dict, client_name: str) -> bytes:
    """
    Parse a QBO AgedReceivableDetail JSON payload and return formatted .xlsx bytes.
    """
    cols_meta  = report.get("Columns", {}).get("Column", [])
    col_titles = [c.get("ColTitle", f"Col{i}") for i, c in enumerate(cols_meta)]
    col_types  = [c.get("ColType", "")          for c in cols_meta]
    n_cols     = len(col_titles)

    wb = Workbook()
    ws = wb.active
    ws.title = "A-R Aging"

    today_str = date.today().strftime("%m/%d/%Y")
    r = 1   # current row cursor

    # ── Title block ───────────────────────────────────────────────────────────
    def _title_row(text, font):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_cols)
        cell = ws.cell(r, 1, value=text)
        cell.font      = font
        cell.alignment = Alignment(horizontal="center")

    _title_row("INCO GROUP INC. — A/R Aging Detail", Font(bold=True, size=14))
    r += 1
    _title_row(client_name, Font(bold=True, size=12))
    r += 1
    _title_row(f"As of {today_str}", Font(italic=True, size=10))
    r += 1
    r += 1  # blank row

    # ── Column header row ─────────────────────────────────────────────────────
    header_row = r
    for ci, title in enumerate(col_titles, start=1):
        cell = ws.cell(r, ci, value=title)
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    ws.freeze_panes = ws.cell(r, 1)  # freeze everything above the first data row

    # ── Sections (aging buckets) ──────────────────────────────────────────────
    # Track money column totals for the grand total row
    grand_totals: dict[int, float] = {}

    for section in report.get("Rows", {}).get("Row", []):
        data_rows = section.get("Rows", {}).get("Row", [])
        data_rows = [dr for dr in data_rows if dr.get("type") == "Data"]
        if not data_rows:
            continue  # skip empty buckets

        # Bucket label (e.g., "Current", "1 - 30")
        bucket_name = (section.get("Header", {}).get("ColData") or [{}])[0].get("value", "")

        # Bucket header row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_cols)
        cell = ws.cell(r, 1, value=bucket_name)
        cell.font  = _BUCKET_FONT
        cell.fill  = _BUCKET_FILL
        cell.alignment = Alignment(horizontal="left", indent=1)
        r += 1

        # Data rows
        for dr in data_rows:
            col_data = dr.get("ColData", [])
            for ci, (cd, ctype) in enumerate(zip(col_data, col_types), start=1):
                raw  = (cd.get("value") or "").strip()
                cell = ws.cell(r, ci)
                cell.font = _NORMAL_FONT
                cell.alignment = Alignment(horizontal="left")

                if ctype in _DATE_COLTYPES:
                    cell.value         = _parse_date(raw)
                    cell.number_format = _DATE_FMT
                elif ctype in _MONEY_COLTYPES:
                    val = _parse_float(raw)
                    cell.value         = val
                    cell.number_format = _MONEY_FMT
                    cell.alignment     = Alignment(horizontal="right")
                    if isinstance(val, float):
                        grand_totals[ci] = grand_totals.get(ci, 0.0) + val
                elif ctype in _INT_COLTYPES:
                    cell.value         = _parse_int(raw)
                    cell.number_format = _INT_FMT
                    cell.alignment     = Alignment(horizontal="right")
                else:
                    cell.value = raw
            r += 1

        # Subtotal row
        summary_data = section.get("Summary", {}).get("ColData", [])
        if summary_data:
            for ci, (cd, ctype) in enumerate(zip(summary_data, col_types), start=1):
                raw  = (cd.get("value") or "").strip()
                cell = ws.cell(r, ci)
                cell.font = _BOLD_FONT
                cell.fill = _SUBTOT_FILL
                if ci == 1:
                    cell.value = f"Total {bucket_name}"
                elif ctype in _MONEY_COLTYPES:
                    cell.value         = _parse_float(raw)
                    cell.number_format = _MONEY_FMT
                    cell.alignment     = Alignment(horizontal="right")
                elif ctype in _DATE_COLTYPES:
                    cell.value         = _parse_date(raw)
                    cell.number_format = _DATE_FMT
                elif ctype in _INT_COLTYPES:
                    cell.value         = _parse_int(raw)
                    cell.number_format = _INT_FMT
                    cell.alignment     = Alignment(horizontal="right")
                else:
                    cell.value = raw
            r += 1

        r += 1  # blank separator between buckets

    # ── Grand total row ───────────────────────────────────────────────────────
    for ci in range(1, n_cols + 1):
        cell = ws.cell(r, ci)
        cell.font   = _BOLD_FONT
        cell.fill   = _GTOTAL_FILL
        cell.border = _TOP_BORDER
        if ci == 1:
            cell.value = "GRAND TOTAL"
        elif ci in grand_totals:
            cell.value         = grand_totals[ci]
            cell.number_format = _MONEY_FMT
            cell.alignment     = Alignment(horizontal="right")

    # ── Auto-size columns ─────────────────────────────────────────────────────
    for ci, title in enumerate(col_titles, start=1):
        col_letter = get_column_letter(ci)
        max_len = len(title)
        for row_cells in ws.iter_rows(
            min_row=header_row + 1, max_row=r,
            min_col=ci, max_col=ci,
        ):
            for cell in row_cells:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 32)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Parsing helpers ───────────────────────────────────────────────────────────

def _parse_date(val: str):
    """Return a Python date for Excel date serialization, or the raw string."""
    for fmt in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(val, fmt).date()
        except ValueError:
            pass
    return val or None


def _parse_float(val: str):
    """Return float or empty string (so the cell stays blank, not 0.00)."""
    try:
        return float(val.replace(",", "")) if val else ""
    except (ValueError, AttributeError):
        return val


def _parse_int(val: str):
    try:
        return int(val) if val else ""
    except (ValueError, AttributeError):
        return val
